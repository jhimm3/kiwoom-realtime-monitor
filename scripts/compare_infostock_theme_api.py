from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from kiwoom_monitor.infrastructure.kiwoom_rest.client import KiwoomRestClient
from kiwoom_monitor.infrastructure.kiwoom_rest.local_config import LocalApiConfig


def _normalized(value: object) -> str:
    return "".join(str(value or "").split()).casefold()


def _load_excel(path: Path) -> dict[str, dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    themes: dict[str, dict[str, object]] = {}
    for theme_value, stock_value, *_rest in sheet.iter_rows(min_row=2, values_only=True):
        theme = str(theme_value or "").strip()
        stock = str(stock_value or "").strip()
        if not theme or not stock:
            continue
        key = _normalized(theme)
        entry = themes.setdefault(key, {"name": theme, "stocks": set()})
        entry["stocks"].add(stock)
    workbook.close()
    return themes


def _load_api(client: KiwoomRestClient) -> dict[str, dict[str, object]]:
    body = {"qry_tp": "0", "stk_cd": "", "date_tp": "10", "thema_nm": "", "flu_pl_amt_tp": "0", "stex_tp": "3"}
    groups: list[dict[str, object]] = []
    cont_yn, next_key = "N", ""
    while True:
        response, has_next, next_key = client.request_with_continuation(
            "ka90001", "/api/dostk/thme", body, cont_yn=cont_yn, next_key=next_key
        )
        groups.extend(value for value in response.get("thema_grp", []) if isinstance(value, dict))
        if not has_next:
            break
        cont_yn = "Y"

    themes: dict[str, dict[str, object]] = {}
    for index, group in enumerate(groups, start=1):
        code = str(group.get("thema_grp_cd", "")).strip()
        name = str(group.get("thema_nm", "")).strip()
        if not code or not name:
            continue
        detail = client.request(
            "ka90002", "/api/dostk/thme",
            {"date_tp": "1", "thema_grp_cd": code, "stex_tp": "3"},
        )
        stocks = {
            str(item.get("stk_nm", "")).strip()
            for item in detail.get("thema_comp_stk", [])
            if isinstance(item, dict) and str(item.get("stk_nm", "")).strip()
        }
        themes[_normalized(name)] = {"name": name, "code": code, "stocks": stocks}
        if index % 25 == 0:
            print(f"Kiwoom groups loaded: {index}/{len(groups)}")
    return themes


def _serializable_theme(entry: dict[str, object]) -> dict[str, object]:
    return {**{key: value for key, value in entry.items() if key != "stocks"}, "stocks": sorted(entry["stocks"])}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--api-config", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args()

    excel = _load_excel(arguments.xlsx)
    api = _load_api(KiwoomRestClient(LocalApiConfig(arguments.api_config).load()))
    common = sorted(excel.keys() & api.keys())
    details = []
    exact_membership = 0
    for key in common:
        excel_stocks = set(excel[key]["stocks"])
        api_stocks = set(api[key]["stocks"])
        if excel_stocks == api_stocks:
            exact_membership += 1
        details.append({
            "theme": str(api[key]["name"]),
            "excel_count": len(excel_stocks),
            "api_count": len(api_stocks),
            "common_count": len(excel_stocks & api_stocks),
            "excel_only_stocks": sorted(excel_stocks - api_stocks),
            "api_only_stocks": sorted(api_stocks - excel_stocks),
        })
    report = {
        "compared_at": datetime.now().isoformat(timespec="seconds"),
        "excel": str(arguments.xlsx),
        "summary": {
            "excel_theme_count": len(excel),
            "excel_mapping_count": sum(len(value["stocks"]) for value in excel.values()),
            "api_theme_count": len(api),
            "api_mapping_count": sum(len(value["stocks"]) for value in api.values()),
            "common_theme_count": len(common),
            "exact_membership_theme_count": exact_membership,
            "excel_only_theme_count": len(excel.keys() - api.keys()),
            "api_only_theme_count": len(api.keys() - excel.keys()),
        },
        "excel_only_themes": sorted(str(excel[key]["name"]) for key in excel.keys() - api.keys()),
        "api_only_themes": sorted(str(api[key]["name"]) for key in api.keys() - excel.keys()),
        "common_theme_details": details,
    }
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
